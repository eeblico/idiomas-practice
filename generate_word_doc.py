import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE

def create_brand_therapy_docx(filename="Cerebro_Digital_Modelo_Cliente.docx"):
    doc = docx.Document()
    
    # Page setup
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # Title
    title = doc.add_paragraph()
    title_run = title.add_run("🧠 Cérebro Digital de Conteúdo")
    title_run.font.name = 'Calibri'
    title_run.font.size = Pt(26)
    title_run.font.bold = True
    title_run.font.color.rgb = RGBColor(126, 34, 206) # Purple

    # Subtitle Callout
    sub = doc.add_paragraph()
    sub_run = sub.add_run("Extração de Contexto & Terapia de Marca (Reunião de 2 Horas)\nCliente: Cliente Exemplo | Nicho: Consultoria Estratégica & Mentorias")
    sub_run.font.name = 'Calibri'
    sub_run.font.size = Pt(11)
    sub_run.font.italic = True
    sub_run.font.color.rgb = RGBColor(100, 116, 139)

    doc.add_paragraph("=" * 60)

    # Section 1
    h1 = doc.add_heading(level=1)
    h1_run = h1.add_run("📖 1. Matéria-Prima de Histórias & Origem")
    h1_run.font.color.rgb = RGBColor(126, 34, 206)
    
    stories = [
        "Passou 5 anos contratando agências tradicionais e recebendo conteúdo genérico sem alma.",
        "Momento de virada: Quando passou a alimentar a IA com causos reais e frases do seu próprio dia a dia.",
        "A revelação de que ninguém lê o rótulo do frasco quando está dentro da garrafa."
    ]
    for s in stories:
        p = doc.add_paragraph(style='List Bullet')
        p.add_run(s)

    # Section 2
    h2 = doc.add_heading(level=1)
    h2_run = h2.add_run("⚡ 2. Crenças Centrais & Teses de Mercado")
    h2_run.font.color.rgb = RGBColor(22, 163, 74)

    beliefs = [
        "Conteúdo genérico de IA é reflexo de falta de contexto, não culpa da ferramenta.",
        "Se o seu posicionamento não incomoda ninguém, ele não atrai ninguém relevante.",
        "A IA é a melhor aceleradora de produção de conteúdo, desde que você forneça a matéria-prima autêntica."
    ]
    for b in beliefs:
        p = doc.add_paragraph(style='List Bullet')
        p.add_run(b)

    # Section 3
    h3 = doc.add_heading(level=1)
    h3_run = h3.add_run("🛡️ 3. Antíteses & O Que Nós Odiamos no Mercado")
    h3_run.font.color.rgb = RGBColor(220, 38, 38)

    antitheses = [
        "Contra o mito de que 'você precisa postar 3 vezes ao day conteúdo genérico'.",
        "Contra agências que cobram caro para entregar cópias sem personalidade.",
        "Contra a ideia de que criar conteúdo autêntico exige horas diárias de escrita do zero."
    ]
    for a in antitheses:
        p = doc.add_paragraph(style='List Bullet')
        p.add_run(a)

    # Section 4
    h4 = doc.add_heading(level=1)
    h4_run = h4.add_run("🗣️ 4. Vocabulário Único & Tom de Voz Autêntico")
    h4_run.font.color.rgb = RGBColor(37, 99, 235)

    vocab = [
        '"Cérebro Digital de Conteúdo"',
        '"Matéria-prima bruta"',
        '"Rótulo dentro da garrafa"',
        '"Provocação estratégica"',
        '"Repertório autêntico"'
    ]
    for v in vocab:
        p = doc.add_paragraph(style='List Bullet')
        r = p.add_run(v)
        r.font.bold = True

    # Section 5
    h5 = doc.add_heading(level=1)
    h5_run = h5.add_run("🎯 5. Posicionamento & Ângulos de Ataque")
    h5_run.font.color.rgb = RGBColor(217, 119, 6)

    pos = [
        "Especialista em transformar vivências individuais em Arquitetura de Contexto para IA.",
        "Serviço exclusivo de Terapia de Marca e extração de contexto em 2 horas de videochamada.",
        "Entrega do Cérebro Digital no Word / Google Docs / Notion pronto para gerar 60+ conteúdos."
    ]
    for p_item in pos:
        p = doc.add_paragraph(style='List Bullet')
        p.add_run(p_item)

    # Section 6: Prompts
    hp = doc.add_heading(level=1)
    hp_run = hp.add_run("🤖 Prompts Mestre para Inteligência Artificial (ChatGPT / Claude / DeepSeek)")
    hp_run.font.color.rgb = RGBColor(126, 34, 206)

    prompts = [
        ("Prompt Mestre: O Rótulo de Dentro da Garrafa [Linha Editorial]", 
         "Aja como o estrategista de conteúdo autêntico do meu negócio. Use os dados do meu Cérebro Digital acima para escrever um post provocativo para o LinkedIn sobre por que a maioria das pessoas falha ao usar IA sem contexto prévio."),
        ("Prompt Mestre: Conversão de Alta Autoridade [Posts de Venda]", 
         "Escreva uma sequência de 3 stories do Instagram contando a história de virada do meu Cérebro Digital. O objetivo é mostrar que uma sessão de 2 horas entrega mais repertório que 6 meses de agência tradicional."),
        ("Carrossel: Os 5 Erros da IA sem Contexto [Carrosséis]", 
         "Crie a estrutura slide por slide de um carrossel de 7 páginas desmistificando o uso genérico do ChatGPT e apresentando o conceito do Segundo Cérebro em 2 Horas.")
    ]

    for title_text, prompt_text in prompts:
        doc.add_heading(title_text, level=2)
        p = doc.add_paragraph()
        r = p.add_run(prompt_text)
        r.font.name = 'Consolas'
        r.font.size = Pt(9.5)
        r.font.color.rgb = RGBColor(51, 65, 85)

    doc.save(filename)
    print(f"Document saved successfully: {filename}")

if __name__ == "__main__":
    create_brand_therapy_docx()
